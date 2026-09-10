"""Extrai da planilha os orcamentos preenchidos e gera os casos de paridade.

Nada aqui e digitado a mao: entradas e valores esperados saem do arquivo.
Os esperados sao os resultados que o proprio Excel gravou em cache
(`data_only=True`), entao a comparacao e contra o que a AA Montagens ve na tela.

Uso:  python tools/extrair_casos.py
Saida: src/teste/casos-planilha.json
"""
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

RAIZ = Path(__file__).resolve().parent.parent
PLANILHA = RAIZ / "IGREJA SAO MIGUEL ARCANJO PORTA PEROLA 2.xlsx"
SAIDA = RAIZ / "src" / "teste" / "casos-planilha.json"

# Colunas da grade, medidas em docs/mapeamento.md secao 4.3
COL_ITEM, COL_DESC, COL_QUANT, COL_UNID, COL_VALOR, COL_TOTAL = 1, 2, 3, 4, 5, 6
PRIMEIRA_LINHA, ULTIMA_LINHA = 17, 27

# Celulas do bloco de totais (secao 4.4)
CELULAS_TOTAIS = {
    "totalDosServicos": "F28",  # =SUM(F18:F27)
    "total": "F32",             # =F28
    "acrescimoNotaFiscal": "F33",  # digitado
    "subTotal": "F34",          # =F32+F33
    "entrada": "F35",           # digitado
    "aPagar": "H35",            # =F34-F35
}


def centavos(valor):
    """Converte o numero do Excel para centavos inteiros, HALF_UP."""
    if valor is None:
        return None
    return int((abs(float(valor)) * 100) + 0.5) * (1 if float(valor) >= 0 else -1)


def mapa_mesclagem(ws):
    """Celula coberta -> (ancora, quantas linhas a mesclagem cobre).

    Precisa cobrir a faixa inteira, nao so a linha-ancora: numa mesclagem
    E18:E21, a linha 19 tambem pertence ao bloco e nao tem valor proprio.
    """
    mapa = {}
    for faixa in ws.merged_cells.ranges:
        ancora = (faixa.min_row, faixa.min_col)
        altura = faixa.max_row - faixa.min_row + 1
        for linha in range(faixa.min_row, faixa.max_row + 1):
            for coluna in range(faixa.min_col, faixa.max_col + 1):
                mapa[(linha, coluna)] = (ancora, altura)
    return mapa


_COPIA = None


def abrir(caminho, data_only):
    """Abre a planilha, contornando o bloqueio de leitura do OneDrive.

    Neste ambiente o Python recebe PermissionError ate para copiar o arquivo,
    enquanto o `cp` do shell le normalmente. Entao a copia sai por ali.
    """
    global _COPIA
    try:
        return openpyxl.load_workbook(caminho, data_only=data_only)
    except PermissionError:
        pass

    if _COPIA is None:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            _COPIA = tmp.name
        resultado = subprocess.run(
            ["cp", str(caminho), _COPIA], capture_output=True, text=True
        )
        if resultado.returncode != 0:
            raise SystemExit(
                "\n".join(
                    [
                        f"nao consegui ler a planilha: {caminho}",
                        "copie o arquivo para fora do OneDrive e passe o caminho:",
                        "  python tools/extrair_casos.py CAMINHO.xlsx",
                    ]
                )
            )
    return openpyxl.load_workbook(_COPIA, data_only=data_only)


def extrair(caminho):
    valores = abrir(caminho, True)["ORÇAMENTO"]
    formulas = abrir(caminho, False)["ORÇAMENTO"]
    mescla = mapa_mesclagem(formulas)

    def ler(linha, coluna):
        ancora, _ = mescla.get((linha, coluna), ((linha, coluna), 1))
        return valores.cell(ancora[0], ancora[1]).value

    def cobertura(linha, coluna):
        """(ancora, altura). Altura > 1 significa celula de bloco mesclado."""
        return mescla.get((linha, coluna), ((linha, coluna), 1))

    secoes = []
    atual = None
    linha = PRIMEIRA_LINHA
    while linha <= ULTIMA_LINHA:
        item = ler(linha, COL_ITEM)
        descricao = ler(linha, COL_DESC)
        rotulo = "" if item is None else str(item).strip()

        # Cabecalho de secao: numero inteiro na coluna ITEM
        if re.fullmatch(r"\d+", rotulo):
            atual = {
                "numeroNaPlanilha": rotulo,
                "titulo": (descricao or "").strip(),
                "linhas": [],
                "precoFechado": None,
                "linhaPlanilha": linha,
            }
            secoes.append(atual)
            linha += 1
            continue

        # Subitem "N.M"
        if re.fullmatch(r"\d+\.\d+", rotulo) and atual is not None:
            quantidade = ler(linha, COL_QUANT)
            unidade = ler(linha, COL_UNID)

            # O preco pode estar mesclado sobre varias linhas: e um bloco fechado.
            ancora_valor, altura = cobertura(linha, COL_VALOR)
            valor_unitario = valores.cell(*ancora_valor).value

            if altura > 1 and atual["precoFechado"] is None:
                atual["precoFechado"] = centavos(valor_unitario)
                atual["alturaDoBloco"] = altura
                atual["ancoraDoBloco"] = f"E{ancora_valor[0]}"

            atual["linhas"].append(
                {
                    "numeroNaPlanilha": rotulo,
                    "descricao": (descricao or "").strip(),
                    "quantidade": None if quantidade is None else float(quantidade),
                    "unidade": None if unidade is None else str(unidade).strip(),
                    # dentro de bloco fechado o valor pertence ao bloco, nao a linha
                    "valorUnitario": None if altura > 1 else centavos(valor_unitario),
                    # num bloco mesclado o total pertence ao bloco, nao a linha
                    "totalNaPlanilha": (
                        None
                        if cobertura(linha, COL_TOTAL)[1] > 1
                        else centavos(ler(linha, COL_TOTAL))
                    ),
                    "linhaPlanilha": linha,
                }
            )
        linha += 1

    esperados = {}
    for nome, celula in CELULAS_TOTAIS.items():
        bruto = valores[celula].value
        esperados[nome] = 0 if bruto is None else centavos(bruto)

    return {
        "arquivo": PLANILHA.name,
        "cliente": (ler(10, 2) or "").strip(),
        "dataEmissao": str(ler(10, 6))[:10],
        "condicoesPagamento": (ler(35, 1) or "").strip(),
        "secoes": secoes,
        "esperados": esperados,
        "formulas": {
            nome: str(formulas[celula].value) for nome, celula in CELULAS_TOTAIS.items()
        },
    }


def main():
    global PLANILHA
    if len(sys.argv) > 1:
        PLANILHA = Path(sys.argv[1])
    if not PLANILHA.exists():
        print(f"planilha nao encontrada: {PLANILHA}")
        return 1
    caso = extrair(PLANILHA)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(
        json.dumps({"casos": [caso]}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"cliente: {caso['cliente']}   emissao: {caso['dataEmissao']}")
    for secao in caso["secoes"]:
        fechado = secao["precoFechado"]
        marca = f"  [preço fechado {fechado / 100:,.2f} sobre {secao.get('alturaDoBloco')} linhas]" if fechado else ""
        print(f"  seção {secao['numeroNaPlanilha']}: {secao['titulo']}{marca}")
        for l in secao["linhas"]:
            v = "—" if l["valorUnitario"] is None else f"{l['valorUnitario'] / 100:,.2f}"
            t = "—" if l["totalNaPlanilha"] is None else f"{l['totalNaPlanilha'] / 100:,.2f}"
            q = "—" if l["quantidade"] is None else f"{l['quantidade']:g}"
            u = l["unidade"] or "—"
            print(f"     {l['numeroNaPlanilha']:5} q={q:>5} {u:>6} valor={v:>12} total={t:>12}")
    print("\n  esperados (cache do proprio Excel):")
    for nome, v in caso["esperados"].items():
        print(f"     {nome:20} {v / 100:>14,.2f}   <- {caso['formulas'][nome]}")
    print(f"\n-> {SAIDA.relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
